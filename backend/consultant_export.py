"""Consultant Timesheet export — reproduces the client-provided XLSX template
(header block, ten-column data grid, SUMMARY/APPROVAL bands) for an arbitrary date
range. `consultant_template.json` at the repo root holds the header defaults and
dropdown option lists, re-read on every request, same idiom as `export_template.json`.

Built fresh with openpyxl on every request rather than filling in a shipped copy of
the client's file: that file pins its 44 data rows to a fixed range (merged summary
cells, data validations, and SUM/COUNTIF formulas all reference exact row numbers),
and openpyxl's insert_rows/delete_rows moves none of those when the row count
changes — any date range other than the one the file shipped with would silently
produce a broken workbook. Computing row positions once here makes every range
correct by construction.
"""

from __future__ import annotations

import io
import json
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

from backend.export import describe_kind
from backend.models import DayEntry, DayKind
from backend.timezone import date_to_str

_CONFIG_PATH = Path(__file__).resolve().parent.parent / "consultant_template.json"

# Locale-independent lookups — datetime.strftime's %a/%b depend on the process
# locale, which is whatever the container image happens to have set. The client's
# file uses English abbreviations regardless of where this runs, so index these
# instead of formatting.
_WEEKDAY_ABBR = ("Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun")
_MONTH_ABBR = (
    "JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC",
)

_COLUMN_HEADERS = [
    "Date", "Day", "Project / Workstream", "Activity Description", "Deliverable / MVP",
    "Category", "Location", "Hours Worked", "Status", "Consultant Remarks",
]
_COLUMN_WIDTHS = {
    "A": 22.36, "B": 11, "C": 22, "D": 32, "E": 25,
    "F": 24, "G": 24.09, "H": 23.82, "I": 16, "J": 28,
}

# Colours/fonts measured from the client's template.
_NAVY = "FF17365D"
_HEADER_BLUE = "FF4472C4"
_LABEL_GREY = "FFE7E6E6"
_VALUE_YELLOW = "FFFFF2CC"
_TOTAL_GREEN = "FFE2F0D9"
_WHITE = "FFFFFFFF"
_BLACK = "FF000000"
_BLUE_TEXT = "FF0000FF"
_FONT_NAME = "Cambria"

_THIN_BOTTOM = Border(bottom=Side(style="thin"))
_CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
_CENTER = Alignment(horizontal="center", vertical="center")


def _fill(rgb: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=rgb)


class ConsultantConfig:
    def __init__(self, raw: dict) -> None:
        self.header_defaults: dict[str, str] = raw.get("header_defaults", {})
        self.field_defaults: dict[str, str] = raw.get("field_defaults", {})
        self.options: dict[str, list[str]] = raw.get("options", {})


def load_consultant_config() -> ConsultantConfig:
    with open(_CONFIG_PATH, encoding="utf-8") as f:
        return ConsultantConfig(json.load(f))


@dataclass(frozen=True)
class ConsultantHeader:
    consultant_name: str
    project_program: str
    vendor_company: str
    technical_lead: str
    pmo_reviewer: str
    period_start: date
    period_end: date


def _has_content(entry: DayEntry | None) -> bool:
    """Whether a day is worth a row on a weekend — a day off, or a work day with
    anything actually logged on it. A bare row some other flow created in passing
    (e.g. an activity typed then deleted) must not add a phantom Saturday."""
    if entry is None:
        return False
    if entry.kind != DayKind.work:
        return True
    return bool(
        entry.hours
        or entry.summary
        or entry.work_text
        or entry.project
        or entry.category
        or entry.status
        or entry.remarks
        or entry.deliverable
        or entry.location
    )


def row_dates(start: date, end: date, entries_by_date: dict[str, DayEntry]) -> list[date]:
    """Every weekday in the range — matching the client's own blank template, which
    pre-lists only Mon-Fri — plus any weekend date that actually has an entry."""
    out: list[date] = []
    d = start
    while d <= end:
        entry = entries_by_date.get(date_to_str(d))
        if d.weekday() < 5 or _has_content(entry):
            out.append(d)
        d += timedelta(days=1)
    return out


def _row_values(d: date, entry: DayEntry | None, config: ConsultantConfig) -> list[Any]:
    weekday = _WEEKDAY_ABBR[d.weekday()]

    if entry is None or (entry.kind == DayKind.work and not _has_content(entry)):
        # A weekday with nothing logged — only date/day/zero hours, same as the
        # client's own blank template. No "MVP"/"Remote" defaults here: writing
        # those into a day nobody worked would fabricate a claim on the sheet.
        return [d, weekday, "", "", "", "", "", 0.0, "", ""]

    if entry.kind != DayKind.work:
        return [d, weekday, "", describe_kind(entry), "", "", "", 0.0, "", ""]

    defaults = config.field_defaults
    return [
        d,
        weekday,
        entry.project or "",
        entry.summary or entry.work_text or "",
        entry.deliverable or defaults.get("deliverable", ""),
        entry.category or "",
        entry.location or defaults.get("location", ""),
        round(entry.hours or 0.0, 2),
        entry.status or "",
        entry.remarks or "",
    ]


def _title_text(start: date, end: date) -> str:
    def fmt(d: date) -> str:
        return f"{d.day:02d} {_MONTH_ABBR[d.month - 1]} {d.year}"

    return f"CONSULTANT TIMESHEET | {fmt(start)} TO {fmt(end)}"


def _label(ws, coord: str, text: str) -> None:
    cell = ws[coord]
    cell.value = text
    cell.font = Font(name=_FONT_NAME, bold=True, size=11, color="FF666666")
    cell.fill = _fill(_LABEL_GREY)
    cell.alignment = _CENTER


def _header_value(ws, coord: str, value: Any, *, center: bool = True, date_fmt: bool = False, color: str = _BLUE_TEXT) -> None:
    cell = ws[coord]
    cell.value = value
    cell.font = Font(name=_FONT_NAME, color=color)
    cell.fill = _fill(_VALUE_YELLOW)
    cell.border = _THIN_BOTTOM
    cell.alignment = Alignment(horizontal="center" if center else None, vertical="center")
    if date_fmt:
        cell.number_format = "dd-mmm-yyyy"


def _band(ws, row: int, text: str) -> None:
    ws.merge_cells(f"A{row}:J{row}")
    cell = ws[f"A{row}"]
    cell.value = text
    cell.font = Font(name=_FONT_NAME, bold=True, color=_WHITE)
    cell.fill = _fill(_NAVY)


def _stat_label(ws, coord: str, text: str) -> None:
    cell = ws[coord]
    cell.value = text
    cell.font = Font(name=_FONT_NAME, bold=True)
    cell.fill = _fill(_LABEL_GREY)
    cell.alignment = _CENTER


def _stat_formula(ws, coord: str, formula: str) -> None:
    cell = ws[coord]
    cell.value = formula
    cell.font = Font(name=_FONT_NAME, bold=True)
    cell.fill = _fill(_TOTAL_GREEN)
    cell.alignment = _CENTER


def to_consultant_xlsx_bytes(
    entries: list[DayEntry], header: ConsultantHeader, config: ConsultantConfig
) -> bytes:
    entries_by_date = {e.date: e for e in entries}
    dates = row_dates(header.period_start, header.period_end, entries_by_date)
    if not dates:
        raise ValueError("No days in that range.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Timesheet"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.freeze_panes = "A8"
    for col, width in _COLUMN_WIDTHS.items():
        ws.column_dimensions[col].width = width

    # Row 1 — title band
    ws.merge_cells("A1:J1")
    ws.row_dimensions[1].height = 30
    title_cell = ws["A1"]
    title_cell.value = _title_text(header.period_start, header.period_end)
    title_cell.font = Font(name=_FONT_NAME, bold=True, size=17, color=_WHITE)
    title_cell.fill = _fill(_NAVY)
    title_cell.alignment = Alignment(horizontal="center")

    # Row 3
    ws.row_dimensions[3].height = 23.5
    _label(ws, "A3", "Consultant Name")
    ws.merge_cells("B3:C3")
    _header_value(ws, "B3", header.consultant_name, center=False)
    _label(ws, "D3", "Vendor Company")
    ws.merge_cells("E3:F3")
    _header_value(ws, "E3", header.vendor_company)
    _label(ws, "G3", "Period Starting")
    _header_value(ws, "H3", header.period_start, date_fmt=True, color=_BLACK)
    _label(ws, "I3", "Period Ending")
    _header_value(ws, "J3", header.period_end, date_fmt=True, color=_BLACK)

    # Row 4
    ws.row_dimensions[4].height = 23.0
    _label(ws, "A4", "Project / Program")
    ws.merge_cells("B4:C4")
    _header_value(ws, "B4", header.project_program, center=False)
    ws.merge_cells("E4:F4")  # blank in the client's file
    _label(ws, "G4", "Technical Lead")
    _header_value(ws, "H4", header.technical_lead)
    _label(ws, "I4", "PMO Reviewer")
    _header_value(ws, "J4", header.pmo_reviewer)

    # Row 7 — column headers
    ws.row_dimensions[7].height = 34.5
    header_font = Font(name=_FONT_NAME, bold=True, size=11, color=_WHITE)
    header_fill = _fill(_HEADER_BLUE)
    for i, text in enumerate(_COLUMN_HEADERS, start=1):
        cell = ws.cell(row=7, column=i, value=text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = _CENTER_WRAP

    # Data rows
    first_row = 8
    for offset, d in enumerate(dates):
        row = first_row + offset
        entry = entries_by_date.get(date_to_str(d))
        values = _row_values(d, entry, config)
        ws.row_dimensions[row].height = 31.5
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col_idx, value=value)
            cell.alignment = _CENTER_WRAP
            cell.border = _THIN_BOTTOM
            col_letter = get_column_letter(col_idx)
            if col_letter == "A":
                cell.number_format = "dd-mmm-yyyy"
                cell.font = Font(name=_FONT_NAME, color=_BLACK)
            elif col_letter == "B":
                cell.font = Font(name=_FONT_NAME, color=_BLACK)
            elif col_letter == "H":
                cell.number_format = "0.00"
                cell.font = Font(name=_FONT_NAME, color=_BLUE_TEXT)
            else:
                cell.font = Font(name=_FONT_NAME, color=_BLUE_TEXT)
    last_row = first_row + len(dates) - 1

    # Data validations — only the client's own option lists/range, over the real span.
    def _list_validation(options: list[str]) -> DataValidation:
        return DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)

    def _span(col: str) -> str:
        return f"{col}{first_row}:{col}{last_row}"

    dv_category = _list_validation(config.options.get("category", []))
    dv_category.add(_span("F"))
    ws.add_data_validation(dv_category)

    dv_status = _list_validation(config.options.get("status", []))
    dv_status.add(_span("I"))
    ws.add_data_validation(dv_status)

    dv_location = _list_validation(config.options.get("location", []))
    dv_location.add(_span("G"))
    ws.add_data_validation(dv_location)

    dv_hours = DataValidation(type="decimal", operator="between", formula1="0", formula2="24", allow_blank=True)
    dv_hours.add(_span("H"))
    ws.add_data_validation(dv_hours)

    # SUMMARY band
    summary_band_row = last_row + 1
    _band(ws, summary_band_row, "SUMMARY")

    totals_row = summary_band_row + 1
    _stat_label(ws, f"A{totals_row}", "Total Hours")
    _stat_formula(ws, f"B{totals_row}", f"=SUM(H{first_row}:H{last_row})")
    _stat_label(ws, f"D{totals_row}", "Days Reported")
    _stat_formula(ws, f"E{totals_row}", f'=COUNTIF(H{first_row}:H{last_row},">0")')
    _stat_label(ws, f"G{totals_row}", "Blocked Items")
    _stat_formula(ws, f"H{totals_row}", f'=COUNTIF(I{first_row}:I{last_row},"Blocked")')

    # APPROVAL band (one blank row after the totals row, matching the client's file)
    approval_band_row = totals_row + 2
    _band(ws, approval_band_row, "APPROVAL")

    approval_row = approval_band_row + 1
    _label(ws, f"A{approval_row}", "Consultant Submission")
    ws.merge_cells(f"B{approval_row}:C{approval_row}")
    ws[f"B{approval_row}"].fill = _fill(_VALUE_YELLOW)
    _label(ws, f"D{approval_row}", "Technical Lead Validation")
    ws.merge_cells(f"E{approval_row}:F{approval_row}")
    ws[f"E{approval_row}"].fill = _fill(_VALUE_YELLOW)
    _label(ws, f"G{approval_row}", "PMO / Budget Validation")
    ws.merge_cells(f"H{approval_row}:I{approval_row}")
    ws[f"H{approval_row}"].fill = _fill(_VALUE_YELLOW)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
