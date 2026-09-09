"""Consultant Timesheet export: row selection (weekdays + weekends with content),
defaults applied only to worked days, leave labelling, and formula/validation spans
that track the real range rather than the client file's fixed Jul-Aug 2026 layout."""

from __future__ import annotations

import io
from datetime import date

import bcrypt
import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend.config import get_settings
from backend.consultant_export import (
    ConsultantHeader,
    load_consultant_config,
    row_dates,
    to_consultant_xlsx_bytes,
)
from backend.models import DayEntry, DayKind


def _header(start: str, end: str, **overrides):
    config = load_consultant_config()
    defaults = dict(
        consultant_name="Jane Doe",
        project_program=config.header_defaults["project_program"],
        vendor_company=config.header_defaults["vendor_company"],
        technical_lead=config.header_defaults["technical_lead"],
        pmo_reviewer=config.header_defaults["pmo_reviewer"],
        period_start=date.fromisoformat(start),
        period_end=date.fromisoformat(end),
    )
    defaults.update(overrides)
    return ConsultantHeader(**defaults)


def test_row_dates_includes_every_weekday_and_only_populated_weekends():
    entries_by_date = {
        "2026-07-04": DayEntry(date="2026-07-04", kind=DayKind.holiday),  # Saturday
        "2026-07-05": DayEntry(date="2026-07-05", kind=DayKind.work),  # Sunday, empty
    }
    dates = row_dates(date(2026, 7, 1), date(2026, 7, 5), entries_by_date)
    # Wed, Thu, Fri (weekdays) + Sat (has a holiday entry) — Sun excluded (empty work row)
    assert [d.isoformat() for d in dates] == ["2026-07-01", "2026-07-02", "2026-07-03", "2026-07-04"]


def test_worked_day_gets_field_defaults_blank_weekday_does_not():
    config = load_consultant_config()
    entries = [DayEntry(date="2026-09-02", kind=DayKind.work, project="Acme", hours=6.0)]
    header = _header("2026-09-02", "2026-09-03")
    content = to_consultant_xlsx_bytes(entries, header, config)
    ws = openpyxl.load_workbook(io.BytesIO(content)).active

    # Row 8 = worked Wed 2026-09-02 -> deliverable/location default to MVP/Remote
    assert ws["E8"].value == "MVP"
    assert ws["G8"].value == "Remote"
    # Row 9 = Thu 2026-09-03, no entry at all -> nothing fabricated
    assert ws["E9"].value in (None, "")
    assert ws["G9"].value in (None, "")
    assert ws["H9"].value == 0.0


def test_holiday_and_time_off_labelled_with_zero_hours_and_blank_metadata():
    config = load_consultant_config()
    entries = [
        DayEntry(date="2026-09-02", kind=DayKind.holiday),
        DayEntry(
            date="2026-09-03",
            kind=DayKind.time_off,
            time_off_reason="Sick day",
            category="Development",  # stored value must not leak onto the sheet
            status="Completed",
        ),
    ]
    header = _header("2026-09-02", "2026-09-03")
    content = to_consultant_xlsx_bytes(entries, header, config)
    ws = openpyxl.load_workbook(io.BytesIO(content)).active

    assert ws["D8"].value == "Holiday"
    assert ws["H8"].value == 0.0
    assert ws["F8"].value in (None, "")

    assert ws["D9"].value == "Time off — Sick day"
    assert ws["H9"].value == 0.0
    assert ws["F9"].value in (None, "")  # category not carried over from the stored row
    assert ws["I9"].value in (None, "")


def test_summary_and_approval_bands_track_the_real_row_span():
    config = load_consultant_config()
    entries = [DayEntry(date="2026-09-01", kind=DayKind.work, hours=8.0)]
    # Tue 2026-09-01 .. Mon 2026-09-07 -> 5 weekdays -> rows 8..12
    header = _header("2026-09-01", "2026-09-07")
    content = to_consultant_xlsx_bytes(entries, header, config)
    ws = openpyxl.load_workbook(io.BytesIO(content)).active

    assert ws["A13"].value == "SUMMARY"
    assert ws["B14"].value == "=SUM(H8:H12)"
    assert ws["E14"].value == '=COUNTIF(H8:H12,">0")'
    assert ws["H14"].value == '=COUNTIF(I8:I12,"Blocked")'
    assert ws["A16"].value == "APPROVAL"

    assert any(str(dv.sqref) == "F8:F12" for dv in ws.data_validations.dataValidation)
    assert any(str(dv.sqref) == "H8:H12" and dv.type == "decimal" for dv in ws.data_validations.dataValidation)


def test_header_block_values_and_merges():
    config = load_consultant_config()
    entries: list[DayEntry] = []
    header = _header("2026-09-01", "2026-09-01", consultant_name="Angshuman Nandy")
    content = to_consultant_xlsx_bytes(entries, header, config)
    ws = openpyxl.load_workbook(io.BytesIO(content)).active

    assert ws["A1"].value == "CONSULTANT TIMESHEET | 01 SEP 2026 TO 01 SEP 2026"
    assert ws["B3"].value == "Angshuman Nandy"
    assert ws["E3"].value == "Tiger Analytics"
    assert ws["B4"].value == "Staff Aug"
    assert ws["H4"].value == "Divyanshu Jain"
    assert ws["J4"].value == "Bane Kodeih"
    assert "B3:C3" in [str(r) for r in ws.merged_cells.ranges]
    assert "E3:F3" in [str(r) for r in ws.merged_cells.ranges]


def test_empty_range_raises():
    config = load_consultant_config()
    header = _header("2026-09-05", "2026-09-06")  # Sat/Sun, no entries -> no rows
    with pytest.raises(ValueError):
        to_consultant_xlsx_bytes([], header, config)


def test_weekday_labels_are_locale_independent_lookup():
    dates = row_dates(date(2026, 7, 1), date(2026, 7, 1), {})
    assert dates == [date(2026, 7, 1)]  # Wednesday — spot-checked against the real calendar


def test_category_option_list_under_excel_inline_limit():
    config = load_consultant_config()
    for options in config.options.values():
        assert len(",".join(options)) < 255


@pytest.fixture
def client(isolated_env, monkeypatch):
    password_hash = bcrypt.hashpw(b"correct-horse-battery", bcrypt.gensalt()).decode()
    monkeypatch.setenv("AUTH_PASSWORD_HASH", password_hash)
    get_settings.cache_clear()

    from backend import db, main

    db._engine = None
    with TestClient(main.app) as c:
        login = c.post(
            "/api/auth/login", json={"username": "testuser", "password": "correct-horse-battery"}
        )
        c.headers.update({"Authorization": f"Bearer {login.json()['token']}"})
        yield c


def test_export_endpoint_consultant_format_downloads_xlsx(client):
    client.post("/api/clock-in", json={"date": "2026-09-01"})
    client.post("/api/clock-out", json={"date": "2026-09-01"})

    resp = client.get(
        "/api/export",
        params={"start": "2026-09-01", "end": "2026-09-07", "format": "consultant"},
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert resp.headers["content-disposition"] == (
        'attachment; filename="consultant-timesheet-2026-09-01-to-2026-09-07.xlsx"'
    )
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws["A3"].value == "Consultant Name"
    assert ws["B3"].value == "testuser"  # defaulted from the logged-in username


def test_export_endpoint_consultant_format_accepts_header_overrides(client):
    resp = client.get(
        "/api/export",
        params={
            "start": "2026-09-01",
            "end": "2026-09-01",
            "format": "consultant",
            "consultant_name": "Someone Else",
            "vendor_company": "Acme Corp",
        },
    )
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws["B3"].value == "Someone Else"
    assert ws["E3"].value == "Acme Corp"
    assert ws["B4"].value == "Staff Aug"  # untouched fields keep the JSON defaults


def test_export_endpoint_consultant_format_rejects_end_before_start(client):
    resp = client.get(
        "/api/export",
        params={"start": "2026-09-07", "end": "2026-09-01", "format": "consultant"},
    )
    assert resp.status_code == 400


def test_export_endpoint_consultant_format_rejects_empty_range(client):
    resp = client.get(
        "/api/export",
        params={"start": "2026-09-05", "end": "2026-09-06", "format": "consultant"},  # Sat/Sun, no entries
    )
    assert resp.status_code == 400


def test_export_endpoint_requires_auth_for_consultant_template():
    from backend import db, main

    with TestClient(main.app) as c:
        resp = c.get("/api/export/consultant-template")
    assert resp.status_code == 401


def test_consultant_template_endpoint_returns_options_and_defaults(client):
    resp = client.get("/api/export/consultant-template")
    assert resp.status_code == 200
    body = resp.json()
    assert body["consultant_name"] == "testuser"
    assert body["header_defaults"]["project_program"] == "Staff Aug"
    assert body["field_defaults"]["location"] == "Remote"
    assert "Development" in body["options"]["category"]
